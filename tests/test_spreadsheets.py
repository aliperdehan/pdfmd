"""v3.29.23: a spreadsheet as the data of a `.csv` block: `::: {.csv file="results.xlsx" sheet="Run 2"}`."""

from __future__ import annotations

import datetime
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import unittest.mock
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pdfmd  # noqa: E402
from pdfmd_tables import sheets  # noqa: E402

try:
    import openpyxl
except ImportError:                                           # pragma: no cover
    openpyxl = None

PANDOC = shutil.which("pandoc")


@unittest.skipIf(openpyxl is None, "needs openpyxl")
class Reading(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        book = openpyxl.Workbook()
        first = book.active
        first.title = "Data"
        first.append(["name", "value", "when"])
        first.append(["a, b", 1.0, datetime.date(2026, 1, 2)])
        first.append(["c", 'x"y', 2.5])
        first.append([None, None, None])
        second = book.create_sheet("Second")
        second.append(["p", "q"])
        second.append([1, 2])
        self.file = self.root / "t.xlsx"
        book.save(self.file)

    def test_the_first_sheet_is_the_default_values_not_formats(self):
        self.assertEqual(sheets.to_csv(self.file), 'name,value,when\n"a, b",1,2026-01-02\nc,"x""y",2.5\n')

    def test_a_sheet_by_name_or_number(self):
        self.assertEqual(sheets.to_csv(self.file, "Second"), "p,q\n1,2\n")
        self.assertEqual(sheets.to_csv(self.file, "2"), "p,q\n1,2\n")

    def test_an_unknown_sheet_names_the_ones_there_are(self):
        with self.assertRaises(sheets.SheetError) as raised:
            sheets.to_csv(self.file, "Nope")
        self.assertIn("Data, Second", str(raised.exception))

    def test_the_csv_is_cached_by_content_and_sheet(self):
        cache = self.root / "cache"
        one = sheets.cached_csv(self.file, None, cache)
        self.assertEqual(one, sheets.cached_csv(self.file, None, cache))
        self.assertNotEqual(one, sheets.cached_csv(self.file, "Second", cache))
        self.assertEqual(one.read_text(encoding="utf-8").splitlines()[0], "name,value,when")

    def test_pandas_reads_it_too(self):
        try:
            import pandas  # noqa: F401
        except ImportError:
            self.skipTest("needs pandas")
        self.assertEqual(sheets._with_pandas(self.file, None), 'name,value,when\n"a, b",1,2026-01-02\nc,"x""y",2.5\n')

    def test_a_file_nothing_can_read_says_what_to_install(self):
        broken = self.root / "b.xlsx"
        broken.write_bytes(b"not a zip")
        with unittest.mock.patch.object(sheets.shutil, "which", lambda name: None):
            with self.assertRaises(sheets.SheetError) as raised:
                sheets.to_csv(broken)
        self.assertIn("Install openpyxl", str(raised.exception))


class Names(unittest.TestCase):
    def test_suffixes(self):
        for name in ("a.xlsx", "A.XLSM", "b.xls", "c.ods"):
            self.assertTrue(sheets.is_spreadsheet(name), name)
        for name in ("a.csv", "b.tsv", "c.xlsx.bak", "d"):
            self.assertFalse(sheets.is_spreadsheet(name), name)


@unittest.skipUnless(PANDOC and openpyxl is not None, "needs Pandoc and openpyxl")
class Build(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        book = openpyxl.Workbook()
        first = book.active
        first.title = "Data"
        first.append(["name", "value"])
        first.append(["alpha", 1.5])
        first.append(["beta", 2])
        second = book.create_sheet("Second")
        second.append(["p", "q"])
        second.append(["x", "y"])
        book.save(self.root / "t.xlsx")
        self.env = {**os.environ, "PDFMD_CONFIG": "", "PDFMD_NO_PROMPT": "1"}

    def run_pdfmd(self, text: str, *args: str) -> subprocess.CompletedProcess:
        (self.root / "d.md").write_text("---\ntitle: T\n---\n\n" + text, encoding="utf-8")
        return subprocess.run([sys.executable, str(ROOT / "pdfmd.py"), "d.md", "--no-stamp", *args], cwd=self.root, env=self.env,
                              capture_output=True, text=True, timeout=300)

    def test_a_sheet_becomes_a_table(self):
        done = self.run_pdfmd('::: {.csv file="t.xlsx"}\n:::\n\n::: {.csv file="t.xlsx" sheet="Second"}\n:::\n', "--to", "gfm", "-o", "o.md")
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        text = (self.root / "o.md").read_text(encoding="utf-8")
        self.assertIn("| alpha | 1.5   |", text)
        self.assertIn("| x   | y   |", text)

    def test_a_missing_spreadsheet_is_a_warning_and_an_empty_block(self):
        done = self.run_pdfmd('::: {.csv file="gone.xlsx"}\n:::\n', "--to", "gfm", "-o", "o.md")
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertIn("gone.xlsx: no such file", done.stderr)
        self.assertIn("is a spreadsheet pdfmd did not convert", done.stdout + done.stderr)

    def test_expand_tables_reads_a_sheet_too(self):
        (self.root / "d.md").write_text('---\ntitle: T\n---\n\n::: {.csv file="t.xlsx"}\n:::\n', encoding="utf-8")
        done = subprocess.run([sys.executable, str(ROOT / "pdfmd.py"), "--expand-tables", "d.md"], cwd=self.root, env=self.env,
                              capture_output=True, text=True, timeout=300)
        self.assertEqual(done.returncode, 0, done.stdout + done.stderr)
        self.assertIn("alpha", (self.root / "d.md").read_text(encoding="utf-8") + done.stdout)


if __name__ == "__main__":
    import unittest.mock  # noqa: F401
    unittest.main()
